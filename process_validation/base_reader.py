"""Lee la base de seguimiento real (.xlsx) de una hoja de periodo. Solo lectura."""

import shutil
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

COL_ID = "ID ESTUDIANTE"
COL_DOCUMENTO = "DOCUMENTO"
COL_ESTADO = "ESTADO"
VALOR_NO_CUMPLE = "NO CUMPLE"
VALOR_OK = "OK"


@dataclass
class BaseSeguimiento:
    # encabezados tal cual la base, en orden, para copiar el diseño en la salida
    encabezados: list
    # id -> {encabezado normalizado: valor}, en el orden de la hoja
    filas: dict
    excluidos_no_cumple: int = 0
    ids_duplicados: list = field(default_factory=list)

    def es_ok(self, id_estudiante):
        fila = self.filas.get(id_estudiante, {})
        return normalizar_texto(fila.get(COL_ESTADO)) == VALOR_OK


def normalizar_texto(valor):
    """Mayúsculas, sin espacios dobles ni saltos de línea. None si está vacío."""
    if valor is None:
        return None
    texto = unicodedata.normalize("NFC", str(valor))
    texto = " ".join(texto.split()).upper()
    return texto or None


def normalizar_id(valor):
    # Excel a veces guarda el ID como número (851049.0)
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


def cargar_base_xlsx(ruta, hoja):
    """Carga la hoja del periodo y descarta los estudiantes con NO CUMPLE."""
    ruta = Path(ruta)
    if not ruta.exists():
        raise FileNotFoundError(f"No se encontró la base en: {ruta}")

    # se lee una copia para no chocar con OneDrive ni con Excel abierto
    with tempfile.TemporaryDirectory() as carpeta_tmp:
        copia = Path(carpeta_tmp) / ruta.name
        shutil.copy2(ruta, copia)
        libro = openpyxl.load_workbook(copia, read_only=True, data_only=True)
        try:
            if hoja not in libro.sheetnames:
                raise ValueError(
                    f"La hoja '{hoja}' no existe. Hojas disponibles: {libro.sheetnames}"
                )
            filas_crudas = list(libro[hoja].iter_rows(values_only=True))
        finally:
            libro.close()

    if not filas_crudas:
        raise ValueError(f"La hoja '{hoja}' está vacía.")

    encabezados = _recortar_encabezados(filas_crudas[0])
    claves = [normalizar_texto(h) for h in encabezados]
    if COL_ID not in claves:
        raise ValueError(f"La hoja '{hoja}' no tiene la columna '{COL_ID}'.")

    base = BaseSeguimiento(encabezados=encabezados, filas={})
    for fila in filas_crudas[1:]:
        registro = {clave: _limpiar_valor(valor) for clave, valor in zip(claves, fila)}
        id_estudiante = normalizar_id(registro.get(COL_ID))
        if id_estudiante is None:
            continue
        if normalizar_texto(registro.get(COL_DOCUMENTO)) == VALOR_NO_CUMPLE:
            base.excluidos_no_cumple += 1
            continue
        if id_estudiante in base.filas:
            # se queda la primera fila; el duplicado se reporta para revisión humana
            base.ids_duplicados.append(id_estudiante)
            continue
        registro[COL_ID] = id_estudiante
        base.filas[id_estudiante] = registro

    return base


def _recortar_encabezados(fila):
    # quita las columnas vacías del final (la hoja 2027-05 trae varias)
    encabezados = [str(h).strip() if h is not None else None for h in fila]
    while encabezados and encabezados[-1] is None:
        encabezados.pop()
    return encabezados


def _limpiar_valor(valor):
    if isinstance(valor, str):
        valor = valor.strip()
        return valor or None
    return valor
