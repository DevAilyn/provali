"""Escribe los resultados con el mismo diseño de la base de seguimiento (ADR-004)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from process_validation.base_reader import COL_ID, normalizar_texto

COL_CARPETA_VACIA = "CARPETA VACÍA"
COL_ESTADO = "ESTADO"
COLUMNAS_SI_NO = ["INDUCCIÓN", "PREINSCRIPCIÓN"]

# requisito del motor -> encabezado normalizado en la base
COLUMNA_REQUISITO = {
    "hoja_vida": "HOJA VIDA",
    "doc_identidad": "DOC IDENTIDAD",
    "afiliacion_salud": "AFIL. SALUD",
    "afiliacion_arl": "AFIL. ARL",
    "carta_presentacion": "CARTA DE PRESENTACIÓN",
    "carta_autorizacion": "CARTA O MEMORANDO DE AUTORIZACIÓN - APROBACIÓN",
    "certificado_laboral": "CERTIFICADO LABORAL CON FUNCIONES",
}

# estado interno -> texto de la base; lo que no está aquí queda vacío para revisión humana
TRADUCCION_ESTADO = {
    "cumplido": "SI",
    "faltante": "NO",
    "no_aplica": "NO APLICA",
    "rechazado_pendiente": "CORREGIR",
}

# mismos colores del formato condicional de la base: (relleno, letra)
COLORES = {
    "SI": ("C6EFCE", "006100"),
    "NO": ("FFC7CE", "9C0006"),
    "NO APLICA": ("CFEBF7", "0F9ED5"),
    "CORREGIR": ("FBE3D6", "9C0006"),
    "OK": ("C6EFCE", "006100"),
    "REVISION": ("FFEB9C", "9C5700"),
}


def traducir_estado(estado):
    return TRADUCCION_ESTADO.get(estado)


def valor_documento(valor_base, estado, carpeta_vacia):
    """Decide qué va en la celda de un documento."""
    # lo que escribió una persona siempre gana
    if valor_base not in (None, ""):
        return valor_base
    traducido = traducir_estado(estado)
    # con carpeta vacía solo se muestra lo que no aplica
    if carpeta_vacia and traducido != "NO APLICA":
        return None
    return traducido


def construir_filas_base(resultados, base):
    """Une la base con los resultados. Devuelve (encabezados, filas, resumen)."""
    claves = base.claves
    pos_id = claves.index(COL_ID)
    encabezados = list(base.encabezados)
    encabezados.insert(pos_id + 1, COL_CARPETA_VACIA)

    por_id = {str(e["id_estudiante"]): e for e in resultados["estudiantes"]}
    resumen = {
        "en_salida": 0,
        "ok_copiados": 0,
        "carpetas_vacias": 0,
        "sin_carpeta": 0,
        "no_cumple_excluidos": base.excluidos_no_cumple,
        "ids_duplicados_en_base": len(base.ids_duplicados),
        "solo_en_carpetas": 0,
    }

    filas = []
    for id_est, registro in base.filas.items():
        est = por_id.get(id_est)
        valores = [registro.get(clave) for clave in claves]
        carpeta = _texto_carpeta(est)

        if base.es_ok(id_est):
            resumen["ok_copiados"] += 1
        elif est is None:
            resumen["sin_carpeta"] += 1
        elif not est.get("error"):
            vacia = est["carpeta_vacia"] is True
            for req, columna in COLUMNA_REQUISITO.items():
                if columna not in claves:
                    continue
                i = claves.index(columna)
                valores[i] = valor_documento(valores[i], est["requisitos"].get(req), vacia)

        if carpeta == "SI":
            resumen["carpetas_vacias"] += 1
        valores.insert(pos_id + 1, carpeta)
        filas.append(valores)

    # estudiantes con carpeta que no están en la base: se agregan al final para no perderlos
    for id_est, est in por_id.items():
        if id_est in base.filas or id_est in base.ids_no_cumple:
            continue
        valores = [None] * len(claves)
        valores[pos_id] = id_est
        if not est.get("error"):
            vacia = est["carpeta_vacia"] is True
            for req, columna in COLUMNA_REQUISITO.items():
                if columna in claves:
                    valores[claves.index(columna)] = valor_documento(
                        None, est["requisitos"].get(req), vacia
                    )
        valores.insert(pos_id + 1, _texto_carpeta(est))
        filas.append(valores)
        resumen["solo_en_carpetas"] += 1

    resumen["en_salida"] = len(filas)
    return encabezados, filas, resumen


def escribir_excel_formato_base(resultados, base, ruta):
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    encabezados, filas, resumen = construir_filas_base(resultados, base)

    libro = Workbook()
    hoja = libro.active
    hoja.title = resultados["periodo"]
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        hoja.append(fila)

    _aplicar_colores(hoja, encabezados, len(filas) + 1)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for i, columna in enumerate(hoja.iter_cols(values_only=True), start=1):
        largo = max((len(str(v)) for v in columna if v is not None), default=10)
        hoja.column_dimensions[get_column_letter(i)].width = min(max(12, largo + 2), 45)

    hoja_resumen = libro.create_sheet("Resumen")
    hoja_resumen.append(["Periodo", resultados["periodo"]])
    hoja_resumen.append(["Generado en", resultados["generado_en"]])
    for clave, valor in resumen.items():
        hoja_resumen.append([clave, valor])
    hoja_resumen.column_dimensions["A"].width = 25

    temporal = ruta.with_suffix(".xlsx.tmp")
    libro.save(temporal)
    temporal.replace(ruta)
    return ruta


def _texto_carpeta(est):
    if est is None or est.get("carpeta_vacia") is None:
        return None
    return "SI" if est["carpeta_vacia"] else "NO"


def _aplicar_colores(hoja, encabezados, ultima_fila):
    # formato condicional, igual que la base: si alguien cambia el texto, el color cambia solo
    claves = [normalizar_texto(h) for h in encabezados]
    columnas_doc = set(COLUMNA_REQUISITO.values()) | set(COLUMNAS_SI_NO) | {"CONTRATO DE APRENDIZAJE"}
    for i, clave in enumerate(claves, start=1):
        letra = get_column_letter(i)
        rango = f"{letra}2:{letra}{max(ultima_fila, 2)}"
        if clave in columnas_doc:
            textos = ["SI", "NO", "NO APLICA", "CORREGIR"]
        elif clave == normalizar_texto(COL_CARPETA_VACIA):
            textos = ["SI", "NO"]
        elif clave == COL_ESTADO:
            textos = ["OK", "REVISION"]
        else:
            continue
        for texto in textos:
            relleno, letra_color = COLORES[texto]
            hoja.conditional_formatting.add(rango, CellIsRule(
                operator="equal",
                formula=[f'"{texto}"'],
                fill=PatternFill(start_color=relleno, end_color=relleno, fill_type="solid"),
                font=Font(color=letra_color),
            ))
