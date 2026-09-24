import json
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from process_validation.evaluator import MATRIZ


def _orden_requisitos() -> list[str]:
    orden = []
    for reglas in MATRIZ.values():
        for requisito in reglas:
            if requisito not in orden:
                orden.append(requisito)
    return orden


ORDEN_REQUISITOS = _orden_requisitos()


def _convertir_si_no(valor: bool | None) -> str | None:
    if valor is None:
        return None
    return "SI" if valor else "NO"


def construir_fila(
    id_estudiante: str,
    total_archivos: int,
    archivos_sin_clasificar: int,
    modalidad: str | None,
    fuente_modalidad: str,
    observacion_modalidad: str | None,
    en_extranjero: bool,
    estados: dict[str, str],
    preinscripcion: bool | None = None,
    induccion: bool | None = None,
) -> dict:

    return {
        "id_estudiante": id_estudiante,
        "carpeta_vacia": total_archivos == 0,
        "modalidad": modalidad,
        "fuente_modalidad": fuente_modalidad,
        "observacion_modalidad": observacion_modalidad or "",
        "en_extranjero": en_extranjero,
        "preinscripcion": _convertir_si_no(preinscripcion),
        "induccion": _convertir_si_no(induccion),
        # un requisito ausente es un fallo del flujo, no un no_aplica
        "requisitos": {req: estados.get(req, "revisar") for req in ORDEN_REQUISITOS},
        "archivos_sin_clasificar": archivos_sin_clasificar,
        "error": None,
    }


def fila_con_error(id_estudiante: str, tipo_error: str) -> dict:
    return {
        "id_estudiante": id_estudiante,
        "carpeta_vacia": None,
        "modalidad": None,
        "fuente_modalidad": "",
        "observacion_modalidad": "",
        "en_extranjero": None,
        "preinscripcion": None,
        "induccion": None,
        "requisitos": {req: "revisar" for req in ORDEN_REQUISITOS},
        "archivos_sin_clasificar": 0,
        "error": tipo_error,
    }

def construir_resultados(periodo: str, filas: list[dict]) -> dict:
    return {
        "periodo": periodo,
        "generado_en": datetime.now().isoformat(timespec="seconds"),
        "resumen": {
            "total_estudiantes": len(filas),
            "carpetas_vacias": sum(f["carpeta_vacia"] is True for f in filas),
            "sin_modalidad": sum(f["modalidad"] is None and not f["error"] for f in filas),
            "con_error": sum(f["error"] is not None for f in filas),
        },
        "estudiantes": filas,
    }


def escribir_json(resultados: dict, ruta: str | Path) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    # se escribe a un temporal y luego se reemplaza, para no dejar un json a medias
    temporal = ruta.with_suffix(".json.tmp")
    with open(temporal, "w", encoding="utf-8") as f:
        json.dump(resultados, f, ensure_ascii=False, indent=2)
    temporal.replace(ruta)
    return ruta

ETIQUETAS_REQUISITOS = {
    "doc_identidad": "Doc. identidad",
    "afiliacion_salud": "Afil. salud",
    "afiliacion_arl": "Afil. ARL",
    "hoja_vida": "Hoja de vida",
    "carta_presentacion": "Carta presentación",
    "carta_autorizacion": "Carta autorización",
    "certificado_laboral": "Certificado laboral",
}

COLORES_ESTADO = {
    "cumplido": "C6EFCE",
    "faltante": "FFC7CE",
    "rechazado_pendiente": "F8CBAD",
    "revisar": "FFEB9C",
    "no_verificable_por_programa": "DDEBF7",
    "no_aplica": "E7E6E6",
}

COLOR_ERROR = "FFC7CE"


def _relleno(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _si_no(valor: bool | None) -> str:
    # None significa sin dato: celda vacía, no un "No"
    if valor is None:
        return ""
    return "Sí" if valor else "No"


def _color_si_no(valor: str | None) -> str | None:
    if valor == "SI":
        return COLORES_ESTADO["cumplido"]
    if valor == "NO":
        return COLORES_ESTADO["faltante"]
    return None


def escribir_excel(resultados: dict, ruta: str | Path) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Estudiantes"

    encabezados = (
        ["ID estudiante", "Carpeta vacía", "Modalidad", "Fuente modalidad",
         "Observación modalidad", "En extranjero", "Preinscripción", "Inducción"]
        + [ETIQUETAS_REQUISITOS.get(req, req) for req in ORDEN_REQUISITOS]
        + ["Archivos sin clasificar", "Error"]
    )
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True)

    col_preinscripcion = 7
    col_induccion = 8
    primera_col_req = 9
    for est in resultados["estudiantes"]:
        hoja.append(
            [est["id_estudiante"], _si_no(est["carpeta_vacia"]), est["modalidad"] or "",
             est["fuente_modalidad"], est["observacion_modalidad"], _si_no(est["en_extranjero"]),
             est["preinscripcion"] or "", est["induccion"] or ""]
            + [est["requisitos"][req] for req in ORDEN_REQUISITOS]
            + [est["archivos_sin_clasificar"], est["error"] or ""]
        )
        fila = hoja.max_row
        for col, valor in ((col_preinscripcion, est["preinscripcion"]), (col_induccion, est["induccion"])):
            color = _color_si_no(valor)
            if color:
                hoja.cell(row=fila, column=col).fill = _relleno(color)
        for i, req in enumerate(ORDEN_REQUISITOS):
            color = COLORES_ESTADO.get(est["requisitos"][req])
            if color:
                hoja.cell(row=fila, column=primera_col_req + i).fill = _relleno(color)
        if est["error"]:
            hoja.cell(row=fila, column=len(encabezados)).fill = _relleno(COLOR_ERROR)

    # encabezado e ID siempre visibles, filtros en todas las columnas
    hoja.freeze_panes = "B2"
    hoja.auto_filter.ref = hoja.dimensions
    # ancho según el texto más largo de cada columna, con tope para observaciones largas
    for i, columna in enumerate(hoja.iter_cols(values_only=True), start=1):
        largo = max(len(str(v)) for v in columna if v is not None)
        hoja.column_dimensions[get_column_letter(i)].width = min(max(12, largo + 2), 45)

    resumen = libro.create_sheet("Resumen")
    resumen.append(["Periodo", resultados["periodo"]])
    resumen.append(["Generado en", resultados["generado_en"]])
    for clave, valor in resultados["resumen"].items():
        resumen.append([clave, valor])
    resumen.column_dimensions["A"].width = 20

    temporal = ruta.with_suffix(".xlsx.tmp")
    libro.save(temporal)
    temporal.replace(ruta)
    return ruta