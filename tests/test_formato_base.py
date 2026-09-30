import openpyxl
import pytest

from process_validation.base_reader import cargar_base_xlsx
from process_validation.formato_base import (
    construir_filas_base,
    escribir_excel_formato_base,
    traducir_estado,
    valor_documento,
)

ENCABEZADOS = [
    "NOMBRES Y APELLIDOS", "ID  ESTUDIANTE", "DOCUMENTO ", "MODALIDAD SELECCIONADA ",
    "INDUCCIÓN", "HOJA VIDA", "DOC \nIDENTIDAD", "AFIL. ARL",
    "CERTIFICADO LABORAL CON FUNCIONES", "OBSERVACIONES", "ESTADO",
]


def _base(tmp_path, filas):
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.title = "2026-23"
    ws.append(ENCABEZADOS)
    for fila in filas:
        ws.append(fila)
    ruta = tmp_path / "base.xlsx"
    libro.save(ruta)
    return cargar_base_xlsx(ruta, "2026-23")


def _est(id_est, carpeta_vacia=False, error=None, **requisitos):
    return {"id_estudiante": id_est, "carpeta_vacia": carpeta_vacia,
            "requisitos": requisitos, "error": error}


def _resultados(*estudiantes):
    return {"periodo": "2026-23", "generado_en": "2026-09-30T17:00:00",
            "estudiantes": list(estudiantes)}


def _fila(encabezados, filas, id_est):
    pos = encabezados.index("ID  ESTUDIANTE")
    fila = next(f for f in filas if f[pos] == id_est)
    return dict(zip(encabezados, fila))


def test_traduccion_de_estados():
    assert traducir_estado("cumplido") == "SI"
    assert traducir_estado("faltante") == "NO"
    assert traducir_estado("no_aplica") == "NO APLICA"
    assert traducir_estado("rechazado_pendiente") == "CORREGIR"
    assert traducir_estado("revisar") is None
    assert traducir_estado("no_verificable_por_programa") is None


def test_valor_humano_gana():
    assert valor_documento("SI", "faltante", carpeta_vacia=False) == "SI"


def test_carpeta_vacia_solo_muestra_no_aplica():
    assert valor_documento(None, "faltante", carpeta_vacia=True) is None
    assert valor_documento(None, "no_aplica", carpeta_vacia=True) == "NO APLICA"


def test_carpeta_vacia_va_junto_al_id(tmp_path):
    base = _base(tmp_path, [["Ana", "100", None, "EMPRENDIMIENTO", "SI", None, None, None, None, None, None]])
    encabezados, _, _ = construir_filas_base(_resultados(_est("100")), base)
    assert encabezados[2] == "CARPETA VACÍA"


def test_llena_documentos_y_respeta_lo_humano(tmp_path):
    base = _base(tmp_path, [["Ana", "100", None, "CONTRATO LABORAL", "SI", "SI", None, None, None, None, "REVISION"]])
    res = _resultados(_est("100", hoja_vida="faltante", doc_identidad="cumplido",
                           afiliacion_arl="revisar", certificado_laboral="rechazado_pendiente"))
    encabezados, filas, _ = construir_filas_base(res, base)
    fila = _fila(encabezados, filas, "100")
    assert fila["HOJA VIDA"] == "SI"
    assert fila["DOC \nIDENTIDAD"] == "SI"
    assert fila["AFIL. ARL"] is None
    assert fila["CERTIFICADO LABORAL CON FUNCIONES"] == "CORREGIR"
    assert fila["CARPETA VACÍA"] == "NO"
    assert fila["MODALIDAD SELECCIONADA"] == "CONTRATO LABORAL"
    assert fila["ESTADO"] == "REVISION"


def test_estado_ok_se_copia_tal_cual(tmp_path):
    base = _base(tmp_path, [["Ana", "100", None, "EMPRENDIMIENTO", "SI", None, None, None, None, "listo", "OK"]])
    res = _resultados(_est("100", hoja_vida="faltante", certificado_laboral="no_aplica"))
    encabezados, filas, resumen = construir_filas_base(res, base)
    fila = _fila(encabezados, filas, "100")
    assert fila["HOJA VIDA"] is None
    assert fila["CERTIFICADO LABORAL CON FUNCIONES"] is None
    assert resumen["ok_copiados"] == 1


def test_carpeta_vacia_deja_documentos_vacios_menos_no_aplica(tmp_path):
    base = _base(tmp_path, [["Ana", "100", None, "EMPRENDIMIENTO", "NO", None, None, None, None, None, None]])
    res = _resultados(_est("100", carpeta_vacia=True, hoja_vida="faltante",
                           certificado_laboral="no_aplica"))
    encabezados, filas, resumen = construir_filas_base(res, base)
    fila = _fila(encabezados, filas, "100")
    assert fila["CARPETA VACÍA"] == "SI"
    assert fila["HOJA VIDA"] is None
    assert fila["CERTIFICADO LABORAL CON FUNCIONES"] == "NO APLICA"
    assert resumen["carpetas_vacias"] == 1


def test_no_cumple_no_sale_aunque_tenga_carpeta(tmp_path):
    base = _base(tmp_path, [["Beto", "200", "NO CUMPLE", None, None, None, None, None, None, None, None]])
    encabezados, filas, _ = construir_filas_base(_resultados(_est("200", hoja_vida="cumplido")), base)
    assert filas == []


def test_estudiante_solo_en_carpetas_se_agrega_al_final(tmp_path):
    base = _base(tmp_path, [["Ana", "100", None, None, None, None, None, None, None, None, None]])
    res = _resultados(_est("100"), _est("300", hoja_vida="cumplido"))
    encabezados, filas, resumen = construir_filas_base(res, base)
    fila = _fila(encabezados, filas, "300")
    assert fila["HOJA VIDA"] == "SI"
    assert fila["NOMBRES Y APELLIDOS"] is None
    assert resumen["solo_en_carpetas"] == 1


def test_escribe_excel_con_formato_condicional(tmp_path):
    base = _base(tmp_path, [["Ana", "100", None, "EMPRENDIMIENTO", "SI", None, None, None, None, None, "OK"]])
    ruta = escribir_excel_formato_base(_resultados(_est("100")), base, tmp_path / "salida" / "resultados.xlsx")
    libro = openpyxl.load_workbook(ruta)
    hoja = libro["2026-23"]
    assert hoja.cell(1, 3).value == "CARPETA VACÍA"
    assert len(hoja.conditional_formatting) > 0
    assert "Resumen" in libro.sheetnames
