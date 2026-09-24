import json

from openpyxl import load_workbook

from process_validation.output import (
    construir_fila,
    fila_con_error,
    construir_resultados,
    escribir_json,
    escribir_excel,
)


def test_construir_fila_rellena_requisito_faltante_como_revisar():
    fila = construir_fila("ST001", 3, 0, "laboral", "csv_real", None, False, {"doc_identidad": "cumplido"})
    assert fila["requisitos"]["doc_identidad"] == "cumplido"
    assert fila["requisitos"]["afiliacion_salud"] == "revisar"


def test_fila_con_error_deja_datos_desconocidos_en_none():
    fila = fila_con_error("ST002", "OSError")
    assert fila["carpeta_vacia"] is None
    assert fila["en_extranjero"] is None
    assert fila["error"] == "OSError"


def test_construir_resultados_cuenta_solo_carpetas_vacias_confirmadas():
    vacia = construir_fila("ST003", 0, 0, None, "sin_resolver", None, False, {})
    con_error = fila_con_error("ST004", "OSError")
    resultados = construir_resultados("2026-99", [vacia, con_error])
    assert resultados["resumen"]["carpetas_vacias"] == 1
    assert resultados["resumen"]["con_error"] == 1


def test_construir_fila_convierte_preinscripcion_induccion():
    fila = construir_fila(
        "ST007", 1, 0, "laboral", "csv_real", None, False, {},
        preinscripcion=True, induccion=False,
    )
    assert fila["preinscripcion"] == "SI"
    assert fila["induccion"] == "NO"


def test_construir_fila_deja_preinscripcion_induccion_en_none_por_defecto():
    fila = construir_fila("ST008", 1, 0, "laboral", "csv_real", None, False, {})
    assert fila["preinscripcion"] is None
    assert fila["induccion"] is None


def test_escribir_json_crea_archivo_legible(tmp_path):
    fila = construir_fila("ST005", 1, 0, "laboral", "csv_real", None, False, {"doc_identidad": "cumplido"})
    resultados = construir_resultados("2026-99", [fila])
    ruta = escribir_json(resultados, tmp_path / "resultados.json")

    with open(ruta, encoding="utf-8") as f:
        leido = json.load(f)
    assert leido["estudiantes"][0]["id_estudiante"] == "ST005"


def test_escribir_json_deja_null_cuando_no_hay_dato(tmp_path):
    fila = construir_fila("ST009", 1, 0, "laboral", "csv_real", None, False, {})
    resultados = construir_resultados("2026-99", [fila])
    ruta = escribir_json(resultados, tmp_path / "resultados.json")

    with open(ruta, encoding="utf-8") as f:
        leido = json.load(f)
    assert leido["estudiantes"][0]["preinscripcion"] is None
    assert leido["estudiantes"][0]["induccion"] is None


def test_escribir_excel_crea_una_fila_por_estudiante(tmp_path):
    fila = construir_fila("ST006", 1, 0, "laboral", "csv_real", None, False, {"doc_identidad": "cumplido"})
    resultados = construir_resultados("2026-99", [fila])
    ruta = escribir_excel(resultados, tmp_path / "resultados.xlsx")

    libro = load_workbook(ruta)
    hoja = libro["Estudiantes"]
    assert hoja.cell(row=2, column=1).value == "ST006"
    assert libro["Resumen"] is not None